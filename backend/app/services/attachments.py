"""Attachment inspection and conversion to model content parts (SPEC section 8.2)."""

from __future__ import annotations

import base64
import csv
import io
import logging
import re
import zipfile
from dataclasses import dataclass
from typing import Any

from app.models import Attachment
from app.router.decision import AttachmentInfo, AttachmentKind

logger = logging.getLogger(__name__)

MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_FILES_PER_MESSAGE = 10
INLINE_TEXT_CHARS = 200_000  # ~50k tokens; longer text files are truncated with a notice
SPREADSHEET_PREVIEW_ROWS = 50
TOKENS_PER_IMAGE = 1_000
TOKENS_PER_PDF_PAGE = 800  # used when no text layer can be read (scans)
CHARS_PER_TOKEN = 4

IMAGE_MIMES = {"image/png", "image/jpeg", "image/webp", "image/gif"}
PDF_MIMES = {"application/pdf"}
SPREADSHEET_MIMES = {
    "text/csv",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.ms-excel",
}
TEXT_EXTENSIONS = {
    ".txt", ".md", ".csv", ".tsv", ".json", ".yaml", ".yml", ".xml", ".html", ".css", ".js",
    ".ts", ".tsx", ".jsx", ".py", ".sql", ".sh", ".java", ".go", ".rs", ".c", ".h", ".cpp",
    ".rb", ".php", ".ini", ".toml", ".log",
}  # fmt: skip


def safe_filename(name: str | None) -> str:
    """Last path segment, no control characters (they would break Content-Disposition)."""
    base = (name or "file").replace("\\", "/").rsplit("/", 1)[-1]
    base = re.sub(r"[\x00-\x1f\x7f]", "", base).strip().strip(".") or "file"
    return base[:200]


class AttachmentError(ValueError):
    pass


def _ext(filename: str) -> str:
    dot = filename.rfind(".")
    return filename[dot:].lower() if dot >= 0 else ""


def sniff_mime(data: bytes, filename: str, declared: str) -> str:
    """Decide the type from content, not the browser's claim (SPEC 8.2: type check)."""
    if data.startswith(b"%PDF-"):
        return "application/pdf"
    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png"
    if data[:3] == b"\xff\xd8\xff":
        return "image/jpeg"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image/webp"
    if data[:6] in (b"GIF87a", b"GIF89a"):
        return "image/gif"
    ext = _ext(filename)
    if data[:4] == b"PK\x03\x04" and ext == ".xlsx":
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if ext in TEXT_EXTENSIONS and _is_text(data):
        return "text/csv" if ext == ".csv" else "text/plain"
    raise AttachmentError(
        f"{filename}: unsupported file type. Upload PDFs, images, spreadsheets or text files."
    )


def _is_text(data: bytes) -> bool:
    sample = data[:8192]
    if b"\x00" in sample:
        return False
    try:
        sample.decode("utf-8")
    except UnicodeDecodeError as err:
        # A multi-byte character cut at the sample boundary is still text.
        return err.start >= len(sample) - 3
    return True


def kind_of(mime: str) -> AttachmentKind:
    if mime in IMAGE_MIMES:
        return "image"
    if mime in PDF_MIMES:
        return "pdf"
    if mime in SPREADSHEET_MIMES:
        return "spreadsheet"
    if mime.startswith("text/"):
        return "text"
    return "other"


@dataclass(frozen=True)
class Inspection:
    mime: str
    page_count: int | None
    est_tokens: int


def inspect(data: bytes, filename: str, declared_mime: str) -> Inspection:
    if len(data) > MAX_FILE_BYTES:
        raise AttachmentError(f"{filename}: larger than {MAX_FILE_BYTES // (1024 * 1024)} MB")
    if not data:
        raise AttachmentError(f"{filename}: the file is empty")
    mime = sniff_mime(data, filename, declared_mime)
    kind = kind_of(mime)
    if kind == "image":
        return Inspection(mime, None, TOKENS_PER_IMAGE)
    if kind == "pdf":
        pages, chars = _pdf_stats(data)
        tokens = max(chars // CHARS_PER_TOKEN, (pages or 1) * TOKENS_PER_PDF_PAGE // 4)
        if chars < 50 * (pages or 1):  # little or no text layer: scanned pages
            tokens = (pages or 1) * TOKENS_PER_PDF_PAGE
        return Inspection(mime, pages, tokens)
    if kind == "spreadsheet" and mime != "text/csv":
        check_xlsx_size(data, filename)
        return Inspection(mime, None, len(spreadsheet_preview(data)) // CHARS_PER_TOKEN)
    text_len = min(len(data), INLINE_TEXT_CHARS)
    return Inspection(mime, None, text_len // CHARS_PER_TOKEN)


def _pdf_stats(data: bytes) -> tuple[int | None, int]:
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        pages = len(reader.pages)
        chars = 0
        for page in reader.pages[:200]:
            chars += len(page.extract_text() or "")
        if pages > 200:
            chars = chars * pages // 200
        return pages, chars
    except Exception:  # malformed PDFs are still sent; the model or parser may cope
        logger.info("could not read PDF structure")
        return None, 0


MAX_XLSX_UNCOMPRESSED = 100 * 1024 * 1024


def check_xlsx_size(data: bytes, filename: str = "spreadsheet") -> None:
    """Refuse zip bombs: openpyxl loads sharedStrings.xml fully even in read-only mode."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            total = sum(i.file_size for i in zf.infolist())
    except zipfile.BadZipFile as err:
        raise AttachmentError(f"{filename}: not a valid Excel file") from err
    if total > MAX_XLSX_UNCOMPRESSED:
        raise AttachmentError(f"{filename}: the spreadsheet is too large when unpacked")


def spreadsheet_preview(data: bytes) -> str:
    from openpyxl import load_workbook

    check_xlsx_size(data)

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = io.StringIO()
    writer = csv.writer(out)
    for ws in wb.worksheets:
        out.write(f"# Sheet: {ws.title}\n")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= SPREADSHEET_PREVIEW_ROWS:
                out.write(f"... (first {SPREADSHEET_PREVIEW_ROWS} rows shown)\n")
                break
            writer.writerow(["" if v is None else v for v in row])
    return out.getvalue()


def info(att: Attachment) -> AttachmentInfo:
    return AttachmentInfo(
        kind=kind_of(att.mime), filename=att.filename, est_tokens=att.est_tokens or 0
    )


def content_parts(att: Attachment, data: bytes, *, sandbox_available: bool) -> list[dict[str, Any]]:
    """OpenAI-style content parts for one attachment."""
    kind = kind_of(att.mime)
    if kind == "image":
        b64 = base64.b64encode(data).decode()
        return [{"type": "image_url", "image_url": {"url": f"data:{att.mime};base64,{b64}"}}]
    if kind == "pdf":
        b64 = base64.b64encode(data).decode()
        return [
            {
                "type": "file",
                "file": {
                    "filename": att.filename,
                    "file_data": f"data:application/pdf;base64,{b64}",
                },
            }
        ]
    mount_note = (
        f"\n(The file is also available to run_python at /data/{att.filename}.)"
        if sandbox_available and kind == "spreadsheet"
        else ""
    )
    if kind == "spreadsheet" and att.mime != "text/csv":
        body = spreadsheet_preview(data)
        return [
            {"type": "text", "text": f"[Attached spreadsheet: {att.filename}]\n{body}{mount_note}"}
        ]
    text = data.decode("utf-8", errors="replace")
    notice = ""
    if len(text) > INLINE_TEXT_CHARS:
        text = text[:INLINE_TEXT_CHARS]
        notice = f"\n[Truncated: only the first {INLINE_TEXT_CHARS:,} characters are included.]"
    return [
        {
            "type": "text",
            "text": f"[Attached file: {att.filename}]\n```\n{text}\n```{notice}{mount_note}",
        }
    ]
